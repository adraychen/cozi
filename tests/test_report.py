import io
import unittest
from copy import deepcopy
from unittest.mock import patch
from docx import Document
from streamlit.testing.v1 import AppTest
from utils.report_state import (ensure_report, add_snapshot, new_block, move_block,
    remove_block, set_draft, restore_text, report_evidence, structured)
from utils.report_export import export_docx, export_pdf, report_nodes
from utils.report_dataset import transformation_history


PAYLOAD = {'title': 'Regression run', 'analysis_type': 'Regression',
    'dataset': {'name': 'Sales'}, 'variables': {'response': 'sales', 'predictors': ['price']},
    'settings': {'confidence': .95}, 'results': {'r_squared': .75},
    'tables': [{'title': 'Coefficients', 'columns': ['Variable', 'Estimate'], 'rows': [['price', -2.5]]}],
    'charts': [], 'warnings': ['Association does not imply causation.'], 'interpretation': 'Original interpretation'}

SCRIPT = '''
import pandas as pd
import streamlit as st
from utils.report_ui import render_report
from utils.report_capture import offer_result, reportable
entry = {'name': 'Sales', 'frame': pd.DataFrame({'sales': [1,2], 'price': [3,4]})}
st.session_state['ws_datasets'] = {'sales': entry}
st.session_state['ws_active'] = 'sales'
screen = st.radio('Screen', ['Analyze', 'Report'], key='screen')
if screen == 'Analyze':
    run = st.number_input('Run', value=1, key='run')
    @reportable
    def result():
        offer_result('Regression', 'Run '+str(run), {'response':'sales'}, {'run':run}, {'r_squared': run/10}, key='test')
    result()
else:
    render_report(entry, {'sales':entry}, 'sales')
'''


class ReportStateTests(unittest.TestCase):
    def test_source_snapshot_order_and_exports_without_report_categories(self):
        from utils.report_state import block_parts
        payload = deepcopy(PAYLOAD)
        payload['content'] = [('Subheading', 'Model summary'), ('Table', payload['tables'][0]),
                              ('Text', 'Original finding'), ('Chart', {'title':'Trend','png':b'image'})]
        report = ensure_report({}); block = add_snapshot(report, payload)
        self.assertEqual([kind for kind, _ in block_parts(block)], ['Subheading','Table','Text','Chart'])
        self.assertEqual(block['text'], '')
        nodes = report_nodes(report)
        self.assertFalse(any(kind == 'Subheading' and value in ('Variables','Method/settings','Key results','Warnings/assumptions','Interpretation') for kind, value in nodes))
        payload['content'][2] = ('Text', 'Rerun finding')
        self.assertIn(('Text','Original finding'), block_parts(block))
        block['included'].remove('Charts')
        self.assertNotIn('Chart', [kind for kind, _ in block_parts(block)])
        document = Document(io.BytesIO(export_docx(report)))
        paragraphs = [p.text for p in document.paragraphs]
        self.assertIn('Model summary', paragraphs)
        self.assertIn('Original finding', paragraphs)
        self.assertNotIn('Variables', paragraphs)
        self.assertNotIn('Method/settings', paragraphs)
        self.assertTrue(export_pdf(report).startswith(b'%PDF'))

    def test_dataset_blocks_reuse_metadata_quality_and_actual_pending_steps(self):
        import pandas as pd
        from utils.variable_metadata import infer_metadata
        from utils.report_dataset import dataset_snapshot
        entry = {'name':'Sales','source':'Upload','frame':pd.DataFrame({'sales':[1.,1.,None], 'group':['a','a','b']})}
        metadata = infer_metadata(entry['frame'])
        overview = dataset_snapshot('Dataset Overview', entry, metadata, {'a':entry}, 'a')
        quality = dataset_snapshot('Data Quality Summary', entry, metadata, {'a':entry}, 'a')
        history = dataset_snapshot('Data Preparation', entry, metadata, {'a':entry}, 'a', {'steps':[{'operation':'Rename column'}]})
        self.assertEqual(overview['results']['Observations'], 3)
        self.assertEqual(quality['results']['Extra duplicate rows'], 1)
        self.assertEqual(history['results']['Performed transformations'][0]['settings']['operation'], 'Rename column')

    def test_multiple_snapshots_are_independent_of_source_and_each_other(self):
        report = ensure_report({})
        source = deepcopy(PAYLOAD)
        first = add_snapshot(report, source)
        source['results']['r_squared'] = .9
        source['tables'][0]['rows'][0][1] = 99
        second = add_snapshot(report, source)
        self.assertEqual(first['snapshot']['results']['r_squared'], .75)
        self.assertEqual(first['snapshot']['tables'][0]['rows'][0][1], -2.5)
        second['snapshot']['variables']['predictors'].append('other')
        self.assertEqual(first['snapshot']['variables']['predictors'], ['price'])
        self.assertNotEqual(first['snapshot']['id'], second['snapshot']['id'])

    def test_order_remove_and_ai_restore(self):
        report = ensure_report({})
        first = add_snapshot(report, PAYLOAD)
        second = new_block('Conclusion'); report['blocks'].append(second)
        move_block(report, second['id'], -1)
        self.assertIs(report['blocks'][0], second)
        set_draft(first, 'Draft'); first['text'] = 'Edited draft'
        set_draft(first, 'Replacement'); restore_text(first)
        self.assertEqual(first['text'], 'Edited draft')
        remove_block(report, second['id']); self.assertEqual(len(report['blocks']), 1)

    def test_evidence_only_uses_included_content_and_user_edits(self):
        report = ensure_report({}); block = add_snapshot(report, PAYLOAD)
        block['included'] = ['Interpretation']; block['text'] = 'Authoritative edit'
        report['blocks'].append(new_block('Executive Summary', snapshot={'results': {'secret': 100}}))
        evidence = str(report_evidence(report))
        self.assertIn('Authoritative edit', evidence)
        self.assertNotIn('r_squared', evidence); self.assertNotIn('secret', evidence)
        self.assertNotIn('Original interpretation', evidence)

    def test_history_is_actual_lineage_and_copied(self):
        registry = {'a': {'name':'Original','source':'Upload'},
            'b': {'name':'Cleaned','source':'Transform','parents':['a'], 'operation':{'steps':[{'operation':'Drop duplicates'}]}}}
        history = transformation_history(registry, 'b')
        self.assertEqual(len(history), 1)
        history[0]['settings']['operation'] = 'changed'
        self.assertEqual(registry['b']['operation']['steps'][0]['operation'], 'Drop duplicates')

    def test_exports_share_edits_tables_order_and_page_breaks(self):
        report = ensure_report({}); report['metadata']['title'] = 'Edited report title'
        first = add_snapshot(report, PAYLOAD); first['text'] = 'Edited narrative'
        first['included'].remove('Warnings/assumptions')
        report['blocks'].extend([new_block('Page Break'), new_block('Conclusion')])
        report['blocks'][-1]['text'] = 'Final conclusion'
        doc = Document(io.BytesIO(export_docx(report)))
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('Edited narrative', text); self.assertNotIn('does not imply', text)
        self.assertLess(text.index('Regression run'), text.index('Conclusion'))
        self.assertEqual(doc.tables[0].cell(1,1).text, '-2.5')
        self.assertIn('w:type="page"', doc.element.xml)
        pdf = export_pdf(report); self.assertTrue(pdf.startswith(b'%PDF'))
        self.assertEqual([kind for kind, _ in report_nodes(report)].count('Page Break'), 1)


class ReportUITests(unittest.TestCase):
    def test_source_output_is_copied_in_order_and_shown_without_generic_sections(self):
        script = SCRIPT.replace("offer_result('Regression'", "from utils import presentation\n        presentation.markdown('**Model summary**')\n        presentation.dataframe(entry['frame'], hide_index=True)\n        presentation.write('Saved finding')\n        offer_result('Regression'")
        app = AppTest.from_string(script).run()
        app.button(key='report_add_test').click().run()
        snapshot = app.session_state['cozy_report']['blocks'][0]['snapshot']
        self.assertEqual([kind for kind, _ in snapshot['content']], ['Subheading','Table','Text'])
        app.radio(key='screen').set_value('Report').run()
        self.assertFalse(app.exception)
        markdown = [m.value for m in app.markdown]
        self.assertIn('**Model summary**', markdown)
        self.assertIn('Saved finding', markdown)
        self.assertNotIn('**Variables**', markdown)
        self.assertNotIn('**Method/settings**', markdown)
        self.assertNotIn('**Key results**', markdown)

    def test_report_renders_saved_chart_with_encoded_numerical_arrays(self):
        import numpy as np
        import plotly.graph_objects as go
        report = ensure_report({})
        payload = deepcopy(PAYLOAD)
        payload['charts'] = [{'title': 'Saved scatter', 'plotly':
            go.Figure(go.Scatter(x=np.arange(10), y=np.arange(10))).to_json()}]
        add_snapshot(report, payload)
        app = AppTest.from_string(SCRIPT)
        app.session_state['screen'] = 'Report'
        app.session_state['cozy_report'] = report
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(len(app.get('plotly_chart')), 1)

    def test_add_reorder_remove_blocks_and_edit_metadata(self):
        app = AppTest.from_string(SCRIPT).run()
        app.radio(key='screen').set_value('Report').run()
        app.text_input(key='report_meta_title').set_value('Custom report').run()
        app.button(key='report_add_section').click().run()
        first = app.session_state['cozy_report']['blocks'][0]['id']
        app.selectbox(key='report_section_type').set_value('Conclusion').run()
        app.button(key='report_add_section').click().run()
        second = app.session_state['cozy_report']['blocks'][1]['id']
        app.button(key='report_up_'+second).click().run()
        self.assertEqual(app.session_state['cozy_report']['blocks'][0]['id'], second)
        app.button(key='report_remove_'+first).click().run()
        self.assertEqual(len(app.session_state['cozy_report']['blocks']), 1)
        self.assertEqual(app.session_state['cozy_report']['metadata']['title'], 'Custom report')
        self.assertFalse(app.exception)

    def test_add_rerun_edit_navigation_and_second_run(self):
        app = AppTest.from_string(SCRIPT).run()
        self.assertFalse(app.button(key='report_add_test').disabled)
        with patch('utils.report_capture.st.toast') as toast:
            app.button(key='report_add_test').click().run()
            toast.assert_not_called()
        self.assertTrue(app.button(key='report_add_test').disabled)
        self.assertEqual(app.button(key='report_add_test').label, '✓ Added to Report')
        app.run()
        self.assertTrue(app.button(key='report_add_test').disabled)
        app.number_input(key='run').set_value(2).run()
        self.assertFalse(app.button(key='report_add_test').disabled)
        self.assertEqual(app.session_state['cozy_report']['blocks'][0]['snapshot']['results']['r_squared'], .1)
        app.button(key='report_add_test').click().run()
        app.radio(key='screen').set_value('Report').run()
        first = app.session_state['cozy_report']['blocks'][0]
        key = 'report_text_'+first['id']
        app.text_area(key=key).set_value('Manual conclusion').run()
        app.radio(key='screen').set_value('Analyze').run()
        app.radio(key='screen').set_value('Report').run()
        self.assertFalse(app.exception)
        self.assertEqual(app.text_area(key=key).value, 'Manual conclusion')
        self.assertEqual(len(app.session_state['cozy_report']['blocks']), 2)

    def test_new_run_identity_reactivates_same_result_and_navigation_keeps_tick(self):
        script = SCRIPT.replace("key='test')", "key='test', run_id=st.session_state.get('new_run_id'))")
        app = AppTest.from_string(script).run()
        app.button(key='report_add_test').click().run()
        app.radio(key='screen').set_value('Report').run()
        app.radio(key='screen').set_value('Analyze').run()
        self.assertTrue(app.button(key='report_add_test').disabled)
        app.session_state['new_run_id'] = 'second training run'
        app.run()
        self.assertFalse(app.button(key='report_add_test').disabled)
        app.button(key='report_add_test').click().run()
        self.assertTrue(app.button(key='report_add_test').disabled)
        self.assertEqual(len(app.session_state['cozy_report']['blocks']), 2)

    def test_ai_regenerate_restore_clear_and_failure_preserve_edits(self):
        app = AppTest.from_string(SCRIPT).run()
        app.button(key='report_add_test').click().run()
        app.radio(key='screen').set_value('Report').run()
        block = app.session_state['cozy_report']['blocks'][0]; bid = block['id']
        app.text_area(key='report_text_'+bid).set_value('Manual text').run()
        with patch('utils.report_ai.generate_report_text', return_value='AI draft') as generate:
            app.button(key='report_ai_'+bid).click().run()
            self.assertIn('Manual text', str(generate.call_args))
        self.assertEqual(app.text_area(key='report_text_'+bid).value, 'AI draft')
        app.button(key='report_undo_'+bid).click().run()
        self.assertEqual(app.text_area(key='report_text_'+bid).value, 'Manual text')
        with patch('utils.report_ai.generate_report_text', side_effect=RuntimeError('offline')):
            app.button(key='report_ai_'+bid).click().run()
        self.assertEqual(app.text_area(key='report_text_'+bid).value, 'Manual text')
        app.button(key='report_clear_'+bid).click().run()
        self.assertEqual(app.text_area(key='report_text_'+bid).value, '')
        app.button(key='report_undo_'+bid).click().run()
        self.assertEqual(app.text_area(key='report_text_'+bid).value, 'Manual text')
        self.assertFalse(app.exception)
