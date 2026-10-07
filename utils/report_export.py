"""DOCX and PDF from the same ordered, included report content."""
from functools import lru_cache
from html import escape
import io
import json
from utils.report_state import block_parts


def cell_text(value):
    if value is None:
        return 'N/A'
    if isinstance(value, float):
        return format(value, '.8g')
    return str(value)


def text_value(value):
    if isinstance(value, dict):
        if 'columns' in value and 'rows' in value:
            return '\n'.join(' | '.join(map(cell_text, row)) for row in [value['columns']]+value['rows'])
        return '\n'.join(str(k)+': '+text_value(v) for k, v in value.items() if v is not None)
    if isinstance(value, list):
        return '\n'.join(text_value(v) for v in value)
    return cell_text(value)


def report_nodes(report):
    metadata = report['metadata']
    nodes = [('Title', metadata['title'])]
    for label, field in [('Dataset', 'dataset'), ('Created', 'created'),
                         ('Observations', 'observations'), ('Variables', 'variables'),
                         ('Author', 'author'), ('Organization', 'organization'), ('Reporting period', 'period')]:
        if metadata.get(field) not in ('', None):
            nodes.append(('Text', f'{label}: {metadata[field]}'))
    for block in report['blocks']:
        if block.get('table_error'):
            raise ValueError(block['table_error'])
        if block['kind'] == 'Page Break':
            nodes.append(('Page Break', None))
            continue
        nodes.append(('Heading', block['title']))
        snapshot = block.get('snapshot') or {}
        if snapshot.get('dataset'):
            nodes.append(('Text', f"{snapshot.get('analysis_type', block['kind'])} · Dataset: {snapshot['dataset']['name']} · Captured: {snapshot.get('timestamp','')}"))
        for kind, value in block_parts(block):
            if kind in ('Text', 'Table', 'Chart'):
                nodes.append((kind, value))
            else:
                nodes.extend([('Subheading', kind), ('Text', text_value(value))])
    return nodes


def table_groups(table):
    columns, rows = table['columns'], table['rows']
    if len(columns) <= 6:
        return [(columns, rows)]
    groups = []
    for start in range(1, len(columns), 5):
        indices = [0]+list(range(start, min(start+5, len(columns))))
        groups.append(([columns[i] for i in indices], [[row[i] for i in indices] for row in rows]))
    return groups


@lru_cache(maxsize=24)
def _plotly_png(configuration):
    import plotly.io as pio
    try:
        return pio.to_image(pio.from_json(configuration), format='png', width=1100, height=650, scale=1.5)
    except Exception as exc:
        raise ValueError('Chart export needs Kaleido and a compatible Chrome/Chromium browser on the server. No chart has been omitted.') from exc


def chart_png(chart):
    return chart['png'] if 'png' in chart else _plotly_png(chart['plotly'])


def export_docx(report):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.8)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Caption'):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(8)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Title'].font.size = Pt(24)
    document.styles['Heading 1'].font.size = Pt(17)
    document.styles['Heading 2'].font.size = Pt(12)
    document.core_properties.title = report['metadata']['title']
    document.core_properties.author = report['metadata']['author']
    for kind, value in report_nodes(report):
        if kind == 'Page Break':
            document.add_page_break()
        elif kind == 'Title':
            document.add_paragraph(value, style='Title')
        elif kind in ('Heading', 'Subheading'):
            document.add_heading(value, level=1 if kind == 'Heading' else 2)
        elif kind == 'Text':
            for paragraph in value.split('\n'):
                document.add_paragraph(paragraph)
        elif kind == 'Chart':
            document.add_paragraph(value['title'], style='Caption').paragraph_format.keep_with_next = True
            png = chart_png(value)
            with Image.open(io.BytesIO(png)) as image:
                width = min(6.9, 6.5*image.width/image.height)
            document.add_picture(io.BytesIO(png), width=Inches(width))
        elif kind == 'Table':
            document.add_paragraph(value.get('title', 'Table'), style='Caption')
            for columns, rows in table_groups(value):
                if not columns:
                    continue
                table = document.add_table(rows=1, cols=len(columns))
                table.style = 'Table Grid'
                table.autofit = False
                weights = [max(6, min(28, max([len(str(c))]+[len(cell_text(r[i])) for r in rows]))) for i, c in enumerate(columns)]
                for i, column in enumerate(columns):
                    table.columns[i].width = Inches(6.9*weights[i]/sum(weights))
                    table.rows[0].cells[i].text = str(column)
                repeat = OxmlElement('w:tblHeader')
                table.rows[0]._tr.get_or_add_trPr().append(repeat)
                for values in rows:
                    cells = table.add_row().cells
                    for cell, content in zip(cells, values):
                        cell.text = cell_text(content)
                for row_index, row in enumerate(table.rows):
                    for column_index, cell in enumerate(row.cells):
                        cell.width = table.columns[column_index].width
                        properties = cell._tc.get_or_add_tcPr()
                        borders = OxmlElement('w:tcBorders')
                        for edge in ('top', 'left', 'bottom', 'right'):
                            line = OxmlElement('w:'+edge)
                            line.set(qn('w:val'), 'single'); line.set(qn('w:sz'), '4'); line.set(qn('w:color'), 'D9D9D9')
                            borders.append(line)
                        properties.append(borders)
                        if row_index == 0:
                            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EDF3'); properties.append(shade)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(4)
                            paragraph.paragraph_format.space_before = Pt(4)
                            for run in paragraph.runs:
                                run.font.size = Pt(10)
                                run.font.bold = row_index == 0
                document.add_paragraph()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(report):
    from pathlib import Path
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
    font_path = Path(reportlab.__file__).parent/'fonts'/'Vera.ttf'
    if 'CozyReport' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('CozyReport', str(font_path)))
    styles = getSampleStyleSheet()
    styles.add(styles['Normal'].clone('Caption'))
    for name in ('Normal', 'Title', 'Heading1', 'Heading2', 'Caption'):
        styles[name].fontName = 'CozyReport'
        styles[name].textColor = colors.black
        styles[name].spaceAfter = 8
    styles['Normal'].fontSize = 11; styles['Normal'].leading = 15
    table_style = styles['Normal'].clone('Table body')
    table_style.fontSize = 9.5; table_style.leading = 12
    def paragraph(text, style='Normal'):
        return Paragraph(escape(str(text)).replace('\n', '<br/>'), styles[style])
    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter, leftMargin=57.6, rightMargin=57.6,
                                 topMargin=57.6, bottomMargin=57.6,
                                 title=report['metadata']['title'], author=report['metadata']['author'])
    story = []
    for kind, value in report_nodes(report):
        if kind == 'Page Break':
            story.append(PageBreak())
        elif kind in ('Title', 'Heading', 'Subheading', 'Text'):
            style = {'Title': 'Title', 'Heading': 'Heading1', 'Subheading': 'Heading2'}.get(kind, 'Normal')
            for line in value.split('\n'):
                story.append(paragraph(line, style))
        elif kind == 'Chart':
            caption = paragraph(value['title'], 'Caption')
            image = Image(io.BytesIO(chart_png(value)))
            scale = min(496.8/image.imageWidth, 440/image.imageHeight)
            image.drawWidth, image.drawHeight = image.imageWidth*scale, image.imageHeight*scale
            story.extend([KeepTogether([caption, image]), Spacer(1, 10)])
        elif kind == 'Table':
            story.append(paragraph(value.get('title', 'Table'), 'Caption'))
            for columns, rows in table_groups(value):
                if not columns:
                    continue
                weights = [max(6, min(28, max([len(str(c))]+[len(cell_text(r[i])) for r in rows]))) for i, c in enumerate(columns)]
                data = [[Paragraph(escape(cell_text(v)), table_style) for v in row] for row in [columns]+rows]
                table = Table(data, colWidths=[496.8*w/sum(weights) for w in weights], repeatRows=1, hAlign='LEFT')
                table.setStyle(TableStyle([('GRID', (0,0), (-1,-1), .5, colors.HexColor('#D9D9D9')),
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E8EDF3')),
                    ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('TOPPADDING', (0,0), (-1,-1), 6),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 6)]))
                story.extend([table, Spacer(1, 10)])
    document.build(story)
    return output.getvalue()
